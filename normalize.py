"""Desempaqueta el extracto xlsx a un CSV listo para COPY INTO.

Tres responsabilidades, en este orden:

1. Validar el contrato del archivo: 21 columnas, en el orden y con los
   encabezados literales que emite la variante de layout de FBL1N. Un cambio
   de variante aguas arriba tiene que romper aca y no seis modelos mas abajo.

2. Validar el checksum contra la fila de totales que SAP agrega al pie. La
   suma corre solo sobre el importe en moneda local: cuando el report mezcla
   monedas, SAP imprime '*' en la columna de moneda de documento porque esa
   columna no es sumable. Si la suma difiere en mas de un centavo, el archivo
   se rechaza entero.

3. Emitir el CSV con los nombres tecnicos de SAP y todos los valores como
   texto en formato de display (fechas DD.MM.YYYY). RAW guarda una imagen fiel
   del origen; el casteo es responsabilidad exclusiva de la capa staging.

La fila de totales nunca llega al CSV: es presentacion, no un hecho.

Uso
---
    python normalize.py                       # todos los extractos sin CSV
    python normalize.py --all                 # rehace todos, existan o no
    python normalize.py --file landing/FBL1N_20260907.xlsx
    python normalize.py --pattern "landing/FBL1N_20260*.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import logging
import sys
from pathlib import Path
from typing import Any

import openpyxl

import config
from generate_extract import extract_date_from_filename

log = logging.getLogger("normalize")


class ContractError(RuntimeError):
    """El archivo no cumple el contrato del extracto."""


def _format_cell(column: str, value: Any) -> str:
    """Convierte una celda del xlsx al texto que va al CSV.

    Todas las columnas viajan como string. Las fechas se escriben en el
    formato de display de SAP (DD.MM.YYYY) para que RAW conserve la misma
    representacion que ve un usuario en la transaccion.
    """
    if value is None or value == "":
        return ""
    if column in config.DATE_COLUMNS:
        if isinstance(value, dt.datetime):
            value = value.date()
        if isinstance(value, dt.date):
            return value.strftime(config.SAP_DATE_FORMAT)
        return str(value).strip()
    if column in config.AMOUNT_COLUMNS:
        return f"{float(value):.2f}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook(path: Path) -> list[tuple[Any, ...]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if config.EXCEL_SHEET_NAME not in workbook.sheetnames:
            raise ContractError(
                f"{path.name}: falta la hoja '{config.EXCEL_SHEET_NAME}' "
                f"(hojas presentes: {workbook.sheetnames})"
            )
        return list(workbook[config.EXCEL_SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()


def validate_header(path: Path, header: tuple[Any, ...]) -> None:
    actual = tuple("" if c is None else str(c).strip() for c in header)
    if actual == config.EXCEL_HEADERS:
        return
    if len(actual) != len(config.EXCEL_HEADERS):
        raise ContractError(
            f"{path.name}: el extracto trae {len(actual)} columnas y el contrato define "
            f"{len(config.EXCEL_HEADERS)}"
        )
    diffs = [
        f"  posicion {i + 1}: esperado {expected!r}, encontrado {found!r}"
        for i, (expected, found) in enumerate(zip(config.EXCEL_HEADERS, actual, strict=True))
        if expected != found
    ]
    raise ContractError(f"{path.name}: encabezados fuera de contrato\n" + "\n".join(diffs))


def split_detail_and_total(
    path: Path, rows: list[tuple[Any, ...]]
) -> tuple[list[tuple[Any, ...]], tuple[Any, ...]]:
    icon = config.COL_INDEX["ICON_STATUS"]
    detail = [r for r in rows if r[icon] != config.TOTAL_ROW_MARKER]
    totals = [r for r in rows if r[icon] == config.TOTAL_ROW_MARKER]

    if len(totals) != 1:
        raise ContractError(
            f"{path.name}: se esperaba exactamente una fila de totales "
            f"({config.TOTAL_ROW_MARKER!r}), se encontraron {len(totals)}"
        )
    if not detail:
        raise ContractError(f"{path.name}: el extracto no tiene filas de detalle")
    return detail, totals[0]


def validate_checksum(path: Path, detail: list[tuple[Any, ...]], total_row: tuple[Any, ...]) -> float:
    """Compara la suma de DMBTR del detalle contra la fila de totales."""
    local = config.COL_INDEX["DMBTR"]
    doc = config.COL_INDEX["WRBTR"]

    declared = total_row[local]
    if declared is None:
        raise ContractError(f"{path.name}: la fila de totales no trae importe en moneda local")

    reported = str(total_row[doc] or "").strip()
    if reported != config.TOTAL_ROW_DOC_AMOUNT:
        # No es fatal: el report imprime un numero cuando todas las partidas
        # comparten moneda. Igual se registra, porque el checksum solo corre
        # sobre moneda local y esa asimetria conviene que quede en el log.
        log.warning(
            "%s: la fila de totales trae %r en moneda de documento en lugar de %r",
            path.name,
            reported,
            config.TOTAL_ROW_DOC_AMOUNT,
        )

    computed = round(sum(float(r[local] or 0.0) for r in detail), 2)
    declared = round(float(declared), 2)
    delta = abs(computed - declared)
    if delta > config.CHECKSUM_TOLERANCE:
        raise ContractError(
            f"{path.name}: checksum roto. Detalle suma {computed:,.2f} y la fila de "
            f"totales declara {declared:,.2f} (delta {delta:,.2f} > "
            f"{config.CHECKSUM_TOLERANCE}). El archivo se rechaza entero."
        )
    log.info(
        "%s | checksum OK: %d filas, %.2f %s (delta %.4f)",
        path.name,
        len(detail),
        declared,
        config.LOCAL_CURRENCY,
        delta,
    )
    return declared


def write_csv(path: Path, detail: list[tuple[Any, ...]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(config.RAW_COLUMNS)
        for row in detail:
            writer.writerow(
                [_format_cell(column, value) for column, value in zip(config.RAW_COLUMNS, row, strict=True)]
            )


def normalize_file(source: Path, staging_dir: Path) -> Path:
    extract_date = extract_date_from_filename(source)
    rows = read_workbook(source)
    if not rows:
        raise ContractError(f"{source.name}: archivo vacio")

    validate_header(source, rows[0])
    detail, total_row = split_detail_and_total(source, rows[1:])
    validate_checksum(source, detail, total_row)

    target = staging_dir / config.STAGING_FILENAME_TEMPLATE.format(
        yyyymmdd=extract_date.strftime("%Y%m%d")
    )
    write_csv(target, detail)
    log.info("%s -> %s (%d filas)", source.name, target.name, len(detail))
    return target


def pending_files(landing_dir: Path, staging_dir: Path, redo_all: bool) -> list[Path]:
    """Extractos a procesar, en orden cronologico."""
    from generate_extract import landing_extracts

    selected = []
    for extract_date, path in landing_extracts(landing_dir):
        csv_path = staging_dir / config.STAGING_FILENAME_TEMPLATE.format(
            yyyymmdd=extract_date.strftime("%Y%m%d")
        )
        if redo_all or not csv_path.exists():
            selected.append(path)
    return selected


def resolve_inputs(args: argparse.Namespace) -> list[Path]:
    if args.file:
        return sorted((Path(f) for f in args.file), key=extract_date_from_filename)
    if args.pattern:
        matches = sorted(Path().glob(args.pattern), key=extract_date_from_filename)
        if not matches:
            raise ContractError(f"El patron {args.pattern!r} no matchea ningun archivo")
        return matches
    return pending_files(config.LANDING_DIR, config.STAGING_DIR, args.all)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Normaliza extractos FBL1N xlsx a CSV para la ingesta."
    )
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--file", action="append", metavar="PATH", help="archivo puntual (repetible)")
    group.add_argument("--pattern", metavar="GLOB", help="glob de extractos a normalizar")
    group.add_argument(
        "--all", action="store_true", help="rehace todos los extractos de landing/"
    )
    parser.add_argument("--verbose", action="store_true", help="log de nivel DEBUG")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(name)s | %(message)s",
    )

    try:
        sources = resolve_inputs(args)
    except (ContractError, ValueError) as exc:
        log.error("%s", exc)
        return 1

    if not sources:
        log.info("No hay extractos pendientes de normalizar en %s", config.LANDING_DIR)
        return 0

    log.info("Normalizando %d extracto(s)", len(sources))
    failures = 0
    for source in sources:
        try:
            normalize_file(source, config.STAGING_DIR)
        except (ContractError, ValueError) as exc:
            # Un archivo roto no frena el backfill entero, pero el proceso
            # termina en error para que el orquestador no siga a la carga.
            log.error("%s", exc)
            failures += 1

    if failures:
        log.error("%d de %d extractos fueron rechazados", failures, len(sources))
        return 1
    log.info("%d extracto(s) normalizados en %s", len(sources), config.STAGING_DIR)
    return 0


if __name__ == "__main__":
    sys.exit(main())
